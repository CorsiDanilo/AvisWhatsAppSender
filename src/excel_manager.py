import os
import re
import shutil
import datetime
from dataclasses import dataclass
from typing import Optional
import openpyxl

def get_desktop_path() -> str:
    """Restituisce il percorso della cartella Desktop dell'utente Windows."""
    try:
        import winreg
        key = winreg.OpenKey(winreg.HKEY_CURRENT_USER, r"Software\Microsoft\Windows\CurrentVersion\Explorer\User Shell Folders")
        desktop_val, _ = winreg.QueryValueEx(key, "Desktop")
        winreg.CloseKey(key)
        expanded = os.path.expandvars(desktop_val)
        if os.path.exists(expanded):
            return expanded
    except Exception:
        pass

    for candidate in [
        os.path.join(os.path.expanduser("~"), "Desktop"),
        os.path.join(os.path.expanduser("~"), "OneDrive", "Desktop"),
        (os.environ.get("USERPROFILE", "") + r"\Desktop") if os.environ.get("USERPROFILE") else ""
    ]:
        if candidate and os.path.exists(candidate):
            return candidate

    return os.path.expanduser("~")

@dataclass
class DonorRecord:
    row_idx: int
    nome: str
    cognome: str
    telefono_raw: str
    telefono_clean: str
    inviato: str  # "Sì", "No", "Saltato"
    is_valid: bool
    note: str = ""

def clean_phone_number(raw: Optional[str]) -> tuple[str, bool]:
    """
    Pulisce e normalizza il numero di telefono per WhatsApp.
    Riconosce numeri italiani a 10 cifre che iniziano per 3, prefissi +39, 0039 o 39.
    Restituisce (numero_formattato, valido).
    """
    if raw is None:
        return ("", False)
    
    raw_str = str(raw).strip()
    if not raw_str:
        return ("", False)

    # Rimuove spazi, trattini, punti, parentesi
    cleaned = re.sub(r"[\s\-\.\(\)\/]", "", raw_str)

    # Se inizia con '+', toglie il '+'
    if cleaned.startswith("+"):
        cleaned = cleaned[1:]

    # Se inizia con '0039', toglie '00'
    if cleaned.startswith("0039"):
        cleaned = cleaned[2:]

    # A questo punto cleaned contiene solo cifre (se era un numero)
    if not cleaned.isdigit():
        return (raw_str, False)

    # Caso 1: inizia con 39 (es. 393331234567)
    if cleaned.startswith("39"):
        radice = cleaned[2:]
        if len(radice) in (9, 10) and radice.startswith("3"):
            return (cleaned, True)
        elif len(radice) >= 8:
            return (cleaned, True)
        return (cleaned, False)

    # Caso 2: numero italiano senza prefisso (10 cifre che inizia per 3, o 9 cifre)
    if len(cleaned) in (9, 10) and cleaned.startswith("3"):
        return (f"39{cleaned}", True)

    # Altri numeri (es. fisso o estero generico se almeno 9 cifre)
    if len(cleaned) >= 9:
        # Se non ha prefisso, per sicurezza in Italia anteponiamo 39
        if not cleaned.startswith("39"):
            return (f"39{cleaned}", True)
        return (cleaned, True)

    return (raw_str, False)


class ExcelManager:
    """Gestisce la lettura e scrittura del file Excel dei donatori."""

    COL_ALIASES_NOME = ["nome", "name", "nominativo"]
    COL_ALIASES_COGNOME = ["cognome", "surname"]
    COL_ALIASES_TELEFONO = ["cellulare", "telefono", "tel", "cell", "mobile", "recapito"]
    COL_ALIASES_STATO = ["stato invio", "stato", "inviato", "esito"]

    def create_desktop_log_copy(self, source_path: str, destination_dir: Optional[str] = None) -> str:
        """
        Crea una copia del file Excel originale sul Desktop con prefisso 'Log_Invio_AVIS_' e timestamp.
        Il file originale non viene mai toccato né modificato.
        Garantisce la presenza della colonna 'Stato Invio' nel nuovo file di log.
        Restituisce il percorso assoluto del file di log creato.
        """
        dest_dir = destination_dir or get_desktop_path()
        if not os.path.exists(dest_dir):
            os.makedirs(dest_dir, exist_ok=True)

        base_name = os.path.splitext(os.path.basename(source_path))[0]
        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        log_filename = f"Log_Invio_AVIS_{base_name}_{timestamp}.xlsx"
        log_path = os.path.join(dest_dir, log_filename)

        # Copia sicura del file sorgente
        shutil.copy2(source_path, log_path)

        # Assicura la presenza della colonna 'Stato Invio' nel file di log
        try:
            wb = openpyxl.load_workbook(log_path)
            ws = wb.active
            if ws is not None:
                header_row = 1
                col_stato = None
                for col in range(1, ws.max_column + 1):
                    val = ws.cell(row=header_row, column=col).value
                    if val is not None and any(alias in str(val).strip().lower() for alias in self.COL_ALIASES_STATO):
                        col_stato = col
                        break
                if col_stato is None:
                    col_stato = ws.max_column + 1
                    ws.cell(row=header_row, column=col_stato).value = "Stato Invio"
                wb.save(log_path)
            wb.close()
        except Exception:
            pass

        return log_path

    def load_file(self, filepath: str) -> list[DonorRecord]:
        """Carica il file Excel e restituisce la lista di DonorRecord."""
        wb = openpyxl.load_workbook(filepath)
        ws = wb.active
        if ws is None:
            wb.close()
            return []

        # Trova gli indici delle colonne (1-based per openpyxl)
        col_nome = None
        col_cognome = None
        col_tel = None
        col_stato = None

        header_row = 1
        for col in range(1, ws.max_column + 1):
            val = ws.cell(row=header_row, column=col).value
            if val is None:
                continue
            val_clean = str(val).strip().lower()

            if col_nome is None and any(alias == val_clean for alias in self.COL_ALIASES_NOME):
                col_nome = col
            elif col_cognome is None and any(alias == val_clean for alias in self.COL_ALIASES_COGNOME):
                col_cognome = col
            elif col_tel is None and any(alias in val_clean for alias in self.COL_ALIASES_TELEFONO):
                col_tel = col
            elif col_stato is None and any(alias in val_clean for alias in self.COL_ALIASES_STATO):
                col_stato = col

        # Se non trova le colonne essenziali, fa una ricerca più flessibile
        if col_tel is None:
            for col in range(1, ws.max_column + 1):
                val = str(ws.cell(row=header_row, column=col).value or "").strip().lower()
                if "tel" in val or "cell" in val or "phone" in val:
                    col_tel = col
                    break

        records: list[DonorRecord] = []

        # Legge le righe a partire dalla riga 2
        for r in range(2, ws.max_row + 1):
            nome_val = ws.cell(row=r, column=col_nome).value if col_nome else ""
            cognome_val = ws.cell(row=r, column=col_cognome).value if col_cognome else ""
            tel_val = ws.cell(row=r, column=col_tel).value if col_tel else ""
            stato_val = ws.cell(row=r, column=col_stato).value if col_stato else ""

            nome_str = str(nome_val).strip() if nome_val is not None else ""
            cognome_str = str(cognome_val).strip() if cognome_val is not None else ""
            tel_str = str(tel_val).strip() if tel_val is not None else ""
            
            # Se la riga è completamente vuota, la ignoriamo
            if not nome_str and not cognome_str and not tel_str:
                continue

            stato_str = str(stato_val).strip() if stato_val is not None else ""
            if not stato_str or stato_str.lower() in ("no", "non inviato"):
                inviato_status = "No"
            elif stato_str.lower().startswith("sì") or stato_str.lower().startswith("si") or "inviato" in stato_str.lower():
                inviato_status = "Sì"
            elif stato_str.lower().startswith("saltato") or stato_str.lower() in ("skip",):
                inviato_status = "Saltato"
            else:
                inviato_status = stato_str

            clean_tel, is_valid = clean_phone_number(tel_str)

            record = DonorRecord(
                row_idx=r,
                nome=nome_str,
                cognome=cognome_str,
                telefono_raw=tel_str,
                telefono_clean=clean_tel,
                inviato=inviato_status,
                is_valid=is_valid
            )
            records.append(record)

        wb.close()
        return records

    def mark_as_sent(self, filepath: str, row_idx: int, status: str = "Sì") -> None:
        """
        Aggiorna la riga row_idx nel file Excel con lo stato indicato (es. 'Sì' o 'Saltato')
        e salva immediatamente il file.
        """
        wb = openpyxl.load_workbook(filepath)
        ws = wb.active
        if ws is None:
            wb.close()
            return

        header_row = 1
        col_stato = None

        for col in range(1, ws.max_column + 1):
            val = ws.cell(row=header_row, column=col).value
            if val is not None and any(alias in str(val).strip().lower() for alias in self.COL_ALIASES_STATO):
                col_stato = col
                break

        # Se la colonna 'Stato Invio' non esiste ancora, la creiamo nell'ultima colonna + 1
        if col_stato is None:
            col_stato = ws.max_column + 1
            ws.cell(row=header_row, column=col_stato).value = "Stato Invio"

        # Aggiorniamo la cella
        now_str = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")
        if status == "Sì":
            ws.cell(row=row_idx, column=col_stato).value = f"Sì ({now_str})"
        else:
            ws.cell(row=row_idx, column=col_stato).value = status

        wb.save(filepath)
        wb.close()

    def get_counts(self, records: list[DonorRecord]) -> dict[str, int]:
        """Restituisce le statistiche su una lista di record."""
        totale = len(records)
        inviati = sum(1 for r in records if r.inviato.startswith("Sì"))
        saltati = sum(1 for r in records if r.inviato.startswith("Saltato"))
        validi = sum(1 for r in records if r.is_valid)
        da_inviare = sum(1 for r in records if r.is_valid and not r.inviato.startswith("Sì") and not r.inviato.startswith("Saltato"))

        return {
            "totale": totale,
            "validi": validi,
            "inviati": inviati,
            "saltati": saltati,
            "da_inviare": da_inviare
        }
